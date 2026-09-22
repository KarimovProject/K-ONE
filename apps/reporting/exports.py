import csv
import io

from django.http import HttpResponse
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.accounts.models import StaffUnavailability


def _event_headers():
    # A function, not a module-level constant: gettext must resolve against
    # the language active for the current request (see ReportExportView's
    # translation.override), not whatever was active at import time.
    return [
        _("Month"),
        _("Date"),
        _("Event"),
        _("Venue"),
        _("Responsible"),
        _("Expected"),
        _("Checked in"),
        _("Attendees"),
        _("Status"),
    ]


def _event_rows(events):
    rows = []
    for event in events.prefetch_related("attendances"):
        attendees = event.attendances.all()
        names = ", ".join(a.attendee_name for a in attendees if a.attendee_name)
        rows.append(
            [
                event.planned_date.strftime("%Y-%m"),
                event.planned_date,
                event.title,
                event.venue.localized_name,
                event.responsible_employee.get_full_name() or event.responsible_employee.username,
                event.expected_attendees,
                len(attendees),
                names,
                event.get_status_display(),
            ]
        )
    return rows


def _busy_staff_rows(start, end):
    slots = (
        StaffUnavailability.objects.filter(start_date__lte=end, end_date__gte=start)
        .select_related("user")
        .order_by("start_date", "start_time")
    )
    return [
        [
            slot.user.get_full_name() or slot.user.username,
            slot.start_date,
            slot.start_time.strftime("%H:%M"),
            slot.end_date,
            slot.end_time.strftime("%H:%M"),
            slot.reason,
        ]
        for slot in slots
    ]


def csv_response(events, report, kind="events"):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="iems-report-events.csv"'
    response.write("\ufeff")
    writer = csv.writer(response)
    datasets = {
        "events": (_event_headers(), _event_rows(events)),
        "venues": (
            [
                _("Venue"),
                _("Available minutes"),
                _("Booked minutes"),
                _("Utilization %"),
                _("Events"),
            ],
            [
                [
                    r["name"],
                    r["available_minutes"],
                    r["booked_minutes"],
                    r["utilization"],
                    r["event_count"],
                ]
                for r in report["venues"]
            ],
        ),
        "attendance": (
            [_("Metric"), _("Value")],
            [[k, v] for k, v in report["attendance"].items() if not isinstance(v, list)],
        ),
        "approvals": (
            [_("Metric"), _("Value")],
            [[k, v] for k, v in report["approvals"].items() if not isinstance(v, dict)],
        ),
        "publications": (
            [_("Metric"), _("Value")],
            [
                [key, value]
                for key, value in report["publications"].items()
                if not isinstance(value, dict | list | tuple | set)
            ],
        ),
    }
    headers, rows = datasets.get(kind, datasets["events"])
    writer.writerow(headers)
    writer.writerows(rows)
    return response


def _append_block(sheet, cursor, title, headers, rows, filterable=False):
    """Write one titled block (title row + header row + data rows) starting
    at 1-indexed row `cursor`. Returns the row number to start the next
    block at (leaves one blank separator row in between)."""
    title_row = cursor
    sheet.cell(row=title_row, column=1, value=title).font = Font(bold=True, size=13)
    header_row = title_row + 1
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True, color="F8FAFC")
        cell.fill = PatternFill("solid", fgColor="07172F")
    for row_offset, row in enumerate(rows, start=1):
        for col, value in enumerate(row, start=1):
            sheet.cell(row=header_row + row_offset, column=col, value=value)
    last_row = header_row + len(rows)
    if filterable and rows:
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
        )
    return last_row + 2  # one blank separator row before the next block


def _autosize_columns(sheet, column_count):
    for column in range(1, column_count + 1):
        letter = get_column_letter(column)
        values = [str(cell.value) for cell in sheet[letter] if cell.value is not None]
        width = max(map(len, values), default=8) + 2
        sheet.column_dimensions[letter].width = min(width, 48)


def xlsx_response(events, report):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _("Report")

    date_range_label = f"{report['start']:%d.%m.%Y} – {report['end']:%d.%m.%Y}"
    sheet.cell(row=1, column=1, value=f"{_('Period')}: {date_range_label}").font = Font(
        bold=True, size=14
    )
    cursor = 3

    cursor = _append_block(
        sheet,
        cursor,
        _("Events in the selected period"),
        _event_headers(),
        _event_rows(events),
        filterable=True,
    )
    cursor = _append_block(
        sheet,
        cursor,
        _("Busy staff"),
        [_("Staff"), _("Start date"), _("Start time"), _("End date"), _("End time"), _("Reason")],
        _busy_staff_rows(report["start"], report["end"]),
    )
    summary_rows = [
        [key.replace("_", " ").title(), value] for key, value in report["summary"].items()
    ]
    cursor = _append_block(sheet, cursor, _("Summary"), [_("Metric"), _("Value")], summary_rows)
    _append_block(
        sheet,
        cursor,
        _("Publications"),
        [_("Metric"), _("Value")],
        [
            [key, value]
            for key, value in report["publications"].items()
            if not isinstance(value, dict | list | tuple | set)
        ],
    )

    max_columns = max(len(_event_headers()), 6)
    _autosize_columns(sheet, max_columns)

    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="iems-report.xlsx"'
    return response


def pdf_response(report):
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph(_("IEMS Leadership Report"), styles["Title"]),
        Paragraph(f"{report['start']:%d.%m.%Y} – {report['end']:%d.%m.%Y}", styles["Normal"]),
        Paragraph(f"{_('Generated')}: {timezone.localtime():%d.%m.%Y %H:%M}", styles["Normal"]),
        Spacer(1, 8 * mm),
    ]
    summary = [[_("Metric"), _("Value")]] + [
        [key.replace("_", " ").title(), "—" if value is None else value]
        for key, value in report["summary"].items()
    ]
    table = Table(summary, colWidths=(90 * mm, 45 * mm), repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#07172F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#F8FAFC")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.extend([table, PageBreak(), Paragraph(_("Venue utilization"), styles["Heading1"])])
    venue_rows = [[_("Venue"), _("Booked minutes"), _("Utilization %"), _("Events")]] + [
        [row["name"], row["booked_minutes"], row["utilization"], row["event_count"]]
        for row in report["venues"]
    ]
    story.append(Table(venue_rows, repeatRows=1))
    for title, section in (
        (_("Attendance"), report["attendance"]),
        (_("Approval performance"), report["approvals"]),
        (_("Publication status"), report["publications"]),
    ):
        story.extend(
            [
                Spacer(1, 7 * mm),
                Paragraph(title, styles["Heading2"]),
                Table(
                    [[_("Metric"), _("Value")]]
                    + [
                        [key.replace("_", " ").title(), "—" if value is None else value]
                        for key, value in section.items()
                        if not isinstance(value, list | dict)
                    ],
                    repeatRows=1,
                ),
            ]
        )
    document.build(story)
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="iems-leadership-report.pdf"'
    return response
