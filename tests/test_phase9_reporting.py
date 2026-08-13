import io
from datetime import time, timedelta

import pytest
from django.contrib.auth import get_user_model
from django.utils import timezone
from openpyxl import load_workbook

from apps.attendance.models import EventAttendance
from apps.audit.models import AuditEventLog
from apps.events.models import Event, EventType
from apps.notifications.models import TelegramConnection, TelegramDelivery
from apps.organizations.models import Organization, Sponsor
from apps.publications.models import Publication
from apps.reporting.analytics import (
    approval_analytics,
    attendance_analytics,
    build_report,
    emergency_analytics,
    event_analytics,
    organization_analytics,
    publication_analytics,
    sponsor_analytics,
    telegram_analytics,
    venue_analytics,
    workload_analytics,
)
from apps.reporting.selectors import filtered_events
from apps.venues.models import Venue

pytestmark = pytest.mark.django_db


@pytest.fixture
def report_data():
    users = get_user_model()
    leader = users.objects.create_user(
        "p9_leader", password="test-password", role=users.Role.LEADERSHIP_VIEWER
    )
    admin = users.objects.create_user(
        "p9_admin", password="test-password", role=users.Role.INTERNATIONAL_ADMIN
    )
    manager = users.objects.create_user(
        "p9_manager", password="test-password", role=users.Role.MANAGEMENT_RESPONSIBLE
    )
    responsible = users.objects.create_user(
        "p9_responsible", password="test-password", role=users.Role.RESPONSIBLE_EMPLOYEE
    )
    reception = users.objects.create_user(
        "p9_reception", password="test-password", role=users.Role.RECEPTION_OPERATOR
    )
    content = users.objects.create_user(
        "p9_content", password="test-password", role=users.Role.CONTENT_MANAGER
    )
    venue = Venue.objects.create(
        code="P9-HALL",
        name_uz="Hisobot zali",
        name_ru="Зал отчётов",
        name_en="Reporting Hall",
        capacity=200,
        working_start=time(8),
        working_end=time(18),
    )
    event_type = EventType.objects.create(
        code="p9-forum", name_uz="Forum", name_ru="Форум", name_en="Forum"
    )
    organization = Organization.objects.create(
        name="P9 International Partner", organization_type=Organization.Type.FOREIGN
    )
    sponsor = Sponsor.objects.create(name="P9 Sponsor")
    today = timezone.localdate()

    def event(title, date, status, start=time(10), end=time(12), expected=100, priority="normal"):
        item = Event.objects.create(
            title=title,
            event_type=event_type,
            venue=venue,
            planned_date=date,
            start_time=start,
            end_time=end,
            responsible_employee=responsible,
            management_responsible=manager,
            created_by=admin,
            status=status,
            expected_attendees=expected,
            priority=priority,
        )
        item.organizing_organizations.add(organization)
        item.sponsors.add(sponsor)
        return item

    approved = event("P9 Approved Forum", today, Event.Status.APPROVED)
    completed = event(
        "P9 Completed Forum", today - timedelta(days=2), Event.Status.COMPLETED, expected=50
    )
    cancelled = event("P9 Cancelled", today - timedelta(days=3), Event.Status.CANCELLED)
    displaced = event("P9 Displaced", today - timedelta(days=1), Event.Status.DISPLACED)
    emergency = event(
        "P9 Emergency",
        today + timedelta(days=1),
        Event.Status.PLANNED,
        priority=Event.Priority.EMERGENCY,
    )
    pending = event("P9 Pending", today, Event.Status.PENDING_APPROVAL)
    now = timezone.now()
    pending.submitted_at = now - timedelta(hours=14)
    pending.save(update_fields=["submitted_at"])
    completed.submitted_at = now - timedelta(hours=8)
    completed.reviewed_at = now - timedelta(hours=2)
    completed.save(update_fields=["submitted_at", "reviewed_at"])
    for action, target in (
        ("event.submitted", completed),
        ("event.approved", completed),
        ("event.rejected", cancelled),
        ("event.emergency_overridden", emergency),
        ("event.displaced", displaced),
    ):
        AuditEventLog.objects.create(action=action, target_id=str(target.pk))
    for index in range(30):
        EventAttendance.objects.create(
            event=approved,
            attendee_identifier_hash=f"p9-{index}",
            attendee_name="" if index < 5 else f"Attendee {index}",
            checkin_method=(
                EventAttendance.Method.PUBLIC_QR
                if index < 20
                else EventAttendance.Method.STAFF_MANUAL
            ),
        )
    publication = Publication.objects.create(
        event=approved,
        platform=Publication.Platform.TELEGRAM_CHANNEL,
        language="uz",
        headline=approved.title,
        created_by=admin,
        status=Publication.Status.PUBLISHED,
        published_at=now,
    )
    Publication.objects.create(
        event=approved,
        platform=Publication.Platform.INSTAGRAM,
        language="en",
        headline=approved.title,
        created_by=admin,
        status=Publication.Status.FAILED,
        retry_count=2,
    )
    TelegramConnection.objects.create(
        user=responsible, chat_id="private-chat-id", telegram_user_id=9911
    )
    TelegramDelivery.objects.create(
        event=approved,
        recipient_user=responsible,
        notification_type="reminder_1d",
        scheduled_for=now,
        status=TelegramDelivery.Status.SENT,
    )
    return {
        "users": (leader, admin, manager, responsible, reception, content),
        "events": (approved, completed, cancelled, displaced, emergency, pending),
        "venue": venue,
        "event_type": event_type,
        "organization": organization,
        "sponsor": sponsor,
        "publication": publication,
        "start": today - timedelta(days=5),
        "end": today + timedelta(days=5),
    }


def queryset(data, user=None, **filters):
    return filtered_events(filters, data["start"], data["end"], user)


def test_executive_kpis_and_previous_comparison(report_data):
    report = build_report(queryset(report_data), report_data["start"], report_data["end"], {})
    assert report["summary"]["total_events"] == 6
    assert report["summary"]["expected_attendees"] == 550
    assert report["summary"]["checked_in"] == 30
    assert report["summary"]["publications_published"] == 1
    assert report["summary"]["venue_utilization"] == 5.5


def test_previous_period_preserves_filters_and_own_event_scope(report_data):
    responsible = report_data["users"][3]
    previous_date = report_data["start"] - timedelta(days=2)
    source = report_data["events"][0]
    values = {
        "event_type": source.event_type,
        "venue": source.venue,
        "planned_date": previous_date,
        "start_time": source.start_time,
        "end_time": source.end_time,
        "management_responsible": source.management_responsible,
        "created_by": source.created_by,
        "status": Event.Status.APPROVED,
        "expected_attendees": 10,
    }
    Event.objects.create(
        title="P9 Previous Own", responsible_employee=responsible, **values
    )
    other = get_user_model().objects.create_user(
        "p9_previous_other", role=get_user_model().Role.RESPONSIBLE_EMPLOYEE
    )
    Event.objects.create(
        title="P9 Previous Other", responsible_employee=other, **values
    )
    query = {
        "period": "custom",
        "start_date": report_data["start"],
        "end_date": report_data["end"],
        "status": Event.Status.APPROVED,
    }
    from apps.reporting.reports import report_context

    _, _, report = report_context(query, responsible)
    assert report["summary"]["previous_total"] == 1


def test_zero_data_has_no_fabricated_percentages(report_data):
    future = report_data["end"] + timedelta(days=100)
    events = filtered_events({}, future, future + timedelta(days=1))
    report = build_report(events, future, future + timedelta(days=1), {})
    assert report["summary"]["attendance_rate"] is None
    assert report["summary"]["reminder_success_rate"] is None


def test_event_groupings_and_filters(report_data):
    events = queryset(report_data, venue=report_data["venue"], status=Event.Status.APPROVED)
    analytics = event_analytics(events)
    assert events.count() == 1
    assert analytics["daily"][0]["total"] == 1
    assert analytics["weekly"] and analytics["monthly"]
    assert analytics["by_status"][0]["label"]
    assert analytics["by_type"][0]["label"] == "Forum"
    assert analytics["by_venue"][0]["total"] == 1
    assert analytics["by_organization"][0]["total"] == 1


def test_venue_working_and_booked_minutes_ignores_invalid_statuses(report_data):
    row = venue_analytics(queryset(report_data), report_data["start"], report_data["end"])[0]
    assert row["available_minutes"] == 11 * 600
    assert row["booked_minutes"] == 360
    assert row["event_count"] == 3


def test_attendance_breakdown(report_data):
    result = attendance_analytics(queryset(report_data))
    assert result["checked"] == 30
    assert result["anonymous"] == 5 and result["identified"] == 25
    assert result["public_qr"] == 20 and result["staff_manual"] == 10


def test_approval_turnaround_and_pending_bucket(report_data):
    result = approval_analytics(queryset(report_data))
    assert result["approved"] == 1 and result["rejected"] == 1
    assert result["average_hours"] == 6.0 and result["median_hours"] == 6.0
    assert result["buckets"]["12_24h"] == 1


def test_emergency_organization_sponsor_and_workload(report_data):
    events = queryset(report_data)
    assert emergency_analytics(events)["emergency_events"] == 1
    assert organization_analytics(events)[0]["total_events"] == 6
    sponsor = sponsor_analytics(events)[0]
    assert sponsor["total_events"] == 6
    assert sponsor["event_types"] == ["Forum"]
    assert len(sponsor["associated_events"]) == 6
    assert workload_analytics(events)[0]["assigned"] == 6


def test_publication_and_telegram_analytics_have_no_identifiers(report_data):
    events = queryset(report_data)
    publication = publication_analytics(events)
    telegram = telegram_analytics(events)
    assert publication["total"] == 2 and publication["retries"] == 2
    assert telegram["sent"] == 1 and telegram["success_rate"] == 100.0
    assert "chat_id" not in str(telegram) and "telegram_user_id" not in str(telegram)


@pytest.mark.parametrize("index", range(6))
def test_reporting_rbac_for_all_relevant_roles(client, report_data, index):
    client.force_login(report_data["users"][index])
    assert client.get("/reports/").status_code == 200


def test_responsible_employee_only_sees_own_events(report_data):
    responsible = report_data["users"][3]
    other = get_user_model().objects.create_user(
        "p9_other", role=get_user_model().Role.RESPONSIBLE_EMPLOYEE
    )
    report_data["events"][0].responsible_employee = other
    report_data["events"][0].save(update_fields=["responsible_employee"])
    assert queryset(report_data, responsible).count() == 5


@pytest.mark.parametrize("kind", ["events", "venues", "attendance", "approvals", "publications"])
def test_csv_exports_are_utf8_and_filtered(client, report_data, kind):
    client.force_login(report_data["users"][1])
    response = client.get(
        "/reports/export/csv/",
        {
            "period": "custom",
            "start_date": report_data["start"],
            "end_date": report_data["end"],
            "kind": kind,
        },
    )
    assert response.status_code == 200 and response.content.startswith(b"\xef\xbb\xbf")
    assert b"attendee_identifier_hash" not in response.content


def test_xlsx_has_professional_sheets(client, report_data):
    client.force_login(report_data["users"][1])
    response = client.get(
        "/reports/export/xlsx/",
        {
            "period": "custom",
            "start_date": report_data["start"],
            "end_date": report_data["end"],
            "language": "en",
        },
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == [
        "Summary",
        "Events",
        "Venues",
        "Attendance",
        "Approvals",
        "Publications",
    ]
    assert workbook["Events"].freeze_panes == "A2"
    assert workbook["Venues"]["A2"].value == "Reporting Hall"


def test_pdf_is_valid_and_nonempty(client, report_data):
    client.force_login(report_data["users"][0])
    response = client.get(
        "/reports/export/pdf/",
        {"period": "custom", "start_date": report_data["start"], "end_date": report_data["end"]},
    )
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF") and len(response.content) > 2_000
    assert response.content.count(b"/Type /Page") >= 2


@pytest.mark.parametrize(
    ("language", "expected"), [("uz", "Hisobot"), ("ru", "Отч"), ("en", "Report")]
)
def test_report_localization_page(client, report_data, language, expected):
    client.force_login(report_data["users"][1])
    client.post("/i18n/setlang/", {"language": language, "next": "/reports/"})
    response = client.get("/reports/")
    assert response.status_code == 200
    assert expected in response.content.decode()


def test_report_api_is_authenticated_and_private(client, report_data):
    assert client.get("/api/v1/reports/summary/").status_code in (403, 302)
    client.force_login(report_data["users"][0])
    response = client.get("/api/v1/reports/summary/")
    text = response.content.decode()
    assert response.status_code == 200
    assert "private-chat-id" not in text and "attendee_identifier_hash" not in text


def test_reception_and_content_report_scope_is_enforced_server_side(client, report_data):
    reception, content = report_data["users"][4], report_data["users"][5]
    client.force_login(reception)
    assert client.get("/api/v1/reports/attendance/").status_code == 200
    assert client.get("/api/v1/reports/publications/").status_code == 403
    assert client.get("/reports/export/csv/?kind=attendance").status_code == 200
    assert client.get("/reports/export/csv/?kind=events").status_code == 403
    assert client.get("/reports/export/xlsx/").status_code == 403
    client.force_login(content)
    assert client.get("/api/v1/reports/publications/").status_code == 200
    assert client.get("/api/v1/reports/attendance/").status_code == 403
    assert client.get("/reports/export/csv/?kind=publications").status_code == 200
